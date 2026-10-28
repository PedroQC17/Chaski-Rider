import ast
import sys
from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent
tree=ast.parse((ROOT/'build_guide.py').read_text(encoding='utf-8'))
data=next(ast.literal_eval(n.value) for n in tree.body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='data' for t in n.targets))
img=Image.open(sys.argv[1]).convert('RGB')
for i,(a,b) in enumerate([(0,226),(226,405),(405,565),(565,785),(785,1006),(1006,1208),(1208,1448)]):
    img.crop((0,int(a*img.height/1448),img.width,int(b*img.height/1448))).save(ROOT/f'color_{i+1}.png')
d=Document()
s=d.sections[0]
s.page_width=Inches(8.27);s.page_height=Inches(11.69)
s.top_margin=Inches(.55);s.bottom_margin=Inches(.55)
s.left_margin=Inches(.66);s.right_margin=Inches(.66)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor.from_string('111111')
    st.paragraph_format.space_after=Pt(7)
d.styles['Normal'].font.size=Pt(10)
d.styles['Normal'].paragraph_format.line_spacing=1.12
d.styles['Title'].font.size=Pt(22)
d.styles['Title'].font.bold=True
d.styles['Heading 1'].font.size=Pt(15)
d.styles['Caption'].font.size=Pt(9)
d.styles['Caption'].font.italic=False
d.styles['Caption'].font.bold=False
d.styles['Caption'].font.color.rgb=RGBColor.from_string('555555')
foot=s.footer.paragraphs[0]
foot.paragraph_format.space_after=Pt(0)
r=foot.add_run('Bosque del Itata  |  Guía de validación manual');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
def p(txt='',style=None,before=0,after=7):
    q=d.add_paragraph(txt,style);q.paragraph_format.space_before=Pt(before);q.paragraph_format.space_after=Pt(after);return q
def labeled(label,txt):
    q=p();q.add_run(label+': ').bold=True;q.add_run(txt)
def header():
    q=p('BOSQUE DEL ITATA',after=8);q.runs[0].bold=True;q.runs[0].font.size=Pt(9)
    p('Guía visual de categorías de vehículos','Title',after=6)
    q=p('Referencia para validación manual',after=14);q.runs[0].font.color.rgb=RGBColor.from_string('555555')
for page,items in enumerate([(0,1),(2,3),(4,5),(6,)]):
    if page:d.add_page_break()
    header()
    if page==0:p('Identifique el tipo de vehículo y revise la presencia de remolque, el número de ejes o la sobredimensión según cada categoría.',after=14)
    for j,i in enumerate(items):
        title,cl,crit,caption,key=data[i]
        q=p(f'Categoría {i+1}',before=17 if j else 0,after=3);q.runs[0].bold=True;q.runs[0].font.size=Pt(9);q.paragraph_format.keep_with_next=True
        p(title,'Heading 1',after=7)
        q=p(after=3);q.add_run().add_picture(str(ROOT/f'color_{i+1}.png'),width=Inches(6.9));q.paragraph_format.keep_with_next=True
        p(caption,'Caption',after=10)
        labeled('Clase de vehículo',cl)
        labeled('Criterio de validación',crit)
        if page==3:labeled('Qué revisar',key)
    if page==3:
        p('Puntos de revisión','Heading 1',before=18,after=10)
        for a,b in [('Tipo de vehículo','Distinguir vehículos livianos, buses, camiones y maquinaria.'),('Remolque','En autos y camionetas, comprobar si arrastran un remolque liviano.'),('Ejes','Comparar dos ejes frente a tres o más en las categorías que lo requieren.'),('Sobredimensión','Revisar esta condición para identificar los vehículos especiales de la categoría 7.')]:labeled(a,b)
        p('Fuente de clasificación  categorias_vehiculos_bosque_del_itata.pdf','Caption',before=20,after=4)
        p('Ilustraciones de referencia para apoyar la identificación visual.','Caption')
for tree in [d.styles.element,d.element]:
    for border in list(tree.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
out=ROOT.parents[1]/'output'/'docx';out.mkdir(parents=True,exist_ok=True)
file=out/'guia_visual_vehiculos_bosque_del_itata.docx';d.save(file);print(file)
